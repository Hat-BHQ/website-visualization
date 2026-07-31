from __future__ import annotations
import io
import json
from html import escape

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import partial
from typing import Any, Callable, Literal
from uuid import UUID

from anyio import to_thread
from fastapi import HTTPException, status

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)
from sqlalchemy.ext.asyncio import AsyncSession

from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

from app.core.config import Settings
from app.repositories.hqa_repository import (
    list_listings,
)
from app.schemas.exports import ListingExportRequest
from app.sync.google_sheets import (
    get_google_credentials,
)


# Giới hạn để tránh một request chiếm toàn bộ RAM.
EXPORT_MAX_ROWS = 50_000

# PDF không phù hợp với số lượng dòng quá lớn.
PDF_MAX_ROWS = 5_000
PDF_MAX_COLUMNS = 10


@dataclass(frozen=True, slots=True)
class ExportColumn:
    key: str
    label: str
    getter: Callable[[Any], Any]


@dataclass(slots=True)
class FileExportResult:
    kind: Literal["file"]
    content: bytes
    filename: str
    media_type: str
    exported_rows: int


@dataclass(slots=True)
class GoogleSheetExportResult:
    kind: Literal["google_sheets"]
    spreadsheet_id: str
    spreadsheet_url: str
    exported_rows: int


def safe_getattr(
    item: Any,
    attribute: str,
) -> Any:
    """
    Một số field chỉ có ở một marketplace.

    Ví dụ:
    - eBay có seller_name.
    - Reverb/Etsy có shop_name.
    - Etsy có shop_id.

    getattr với giá trị mặc định giúp export không bị lỗi.
    """
    return getattr(item, attribute, None)


def get_seller(item: Any) -> str | None:
    return safe_getattr(item, "seller_name") or safe_getattr(item, "shop_name")


EXPORT_COLUMNS: dict[str, ExportColumn] = {
    "id": ExportColumn(
        "id",
        "ID nội bộ",
        lambda item: safe_getattr(item, "id"),
    ),
    "external_listing_id": ExportColumn(
        "external_listing_id",
        "Listing ID",
        lambda item: safe_getattr(
            item,
            "external_listing_id",
        ),
    ),
    "listing_title": ExportColumn(
        "listing_title",
        "Tên sản phẩm",
        lambda item: safe_getattr(
            item,
            "listing_title",
        ),
    ),
    "listing_url": ExportColumn(
        "listing_url",
        "Listing URL",
        lambda item: safe_getattr(
            item,
            "listing_url",
        ),
    ),
    # Alias tương thích cho giao diện cũ. Nếu cần đúng tên cột DB,
    # frontend nên chọn seller_name hoặc shop_name.
    "seller": ExportColumn(
        "seller",
        "seller",
        get_seller,
    ),
    "seller_name": ExportColumn(
        "seller_name",
        "seller_name",
        lambda item: safe_getattr(item, "seller_name"),
    ),
    "shop_name": ExportColumn(
        "shop_name",
        "shop_name",
        lambda item: safe_getattr(item, "shop_name"),
    ),
    "shop_id": ExportColumn(
        "shop_id",
        "Shop ID",
        lambda item: safe_getattr(
            item,
            "shop_id",
        ),
    ),
    "published_at": ExportColumn(
        "published_at",
        "Ngày đăng",
        lambda item: safe_getattr(
            item,
            "published_at",
        ),
    ),
    "listing_location": ExportColumn(
        "listing_location",
        "Vị trí",
        lambda item: safe_getattr(
            item,
            "listing_location",
        ),
    ),
    "country_code": ExportColumn(
        "country_code",
        "Mã quốc gia",
        lambda item: safe_getattr(
            item,
            "country_code",
        ),
    ),
    "category_id": ExportColumn(
        "category_id",
        "Category ID",
        lambda item: safe_getattr(
            item,
            "category_id",
        ),
    ),
    "category_name": ExportColumn(
        "category_name",
        "Danh mục",
        lambda item: safe_getattr(
            item,
            "category_name",
        ),
    ),
    "condition_id": ExportColumn(
        "condition_id",
        "Condition ID",
        lambda item: safe_getattr(
            item,
            "condition_id",
        ),
    ),
    "condition_name": ExportColumn(
        "condition_name",
        "Tình trạng",
        lambda item: safe_getattr(
            item,
            "condition_name",
        ),
    ),
    "image_url": ExportColumn(
        "image_url",
        "Image URL",
        lambda item: safe_getattr(
            item,
            "image_url",
        ),
    ),
    "current_price": ExportColumn(
        "current_price",
        "Giá hiện tại",
        lambda item: safe_getattr(
            item,
            "current_price",
        ),
    ),
    "shipping_price": ExportColumn(
        "shipping_price",
        "Phí vận chuyển",
        lambda item: safe_getattr(
            item,
            "shipping_price",
        ),
    ),
    "total_price": ExportColumn(
        "total_price",
        "Tổng giá",
        lambda item: safe_getattr(
            item,
            "total_price",
        ),
    ),
    "currency": ExportColumn(
        "currency",
        "Tiền tệ",
        lambda item: safe_getattr(
            item,
            "currency",
        ),
    ),
    "quantity": ExportColumn(
        "quantity",
        "quantity",
        lambda item: safe_getattr(item, "quantity"),
    ),
    "listing_views": ExportColumn(
        "listing_views",
        "listing_views",
        lambda item: safe_getattr(
            item,
            "listing_views",
        ),
    ),
    "listing_status": ExportColumn(
        "listing_status",
        "Trạng thái",
        lambda item: safe_getattr(
            item,
            "listing_status",
        ),
    ),
    "status_reason": ExportColumn(
        "status_reason",
        "Lý do trạng thái",
        lambda item: safe_getattr(
            item,
            "status_reason",
        ),
    ),
    "first_seen_at": ExportColumn(
        "first_seen_at",
        "Lần đầu ghi nhận",
        lambda item: safe_getattr(
            item,
            "first_seen_at",
        ),
    ),
    "last_seen_at": ExportColumn(
        "last_seen_at",
        "Cập nhật gần nhất",
        lambda item: safe_getattr(
            item,
            "last_seen_at",
        ),
    ),
    "state_hash": ExportColumn(
        "state_hash",
        "State hash",
        lambda item: safe_getattr(
            item,
            "state_hash",
        ),
    ),
    "created_at": ExportColumn(
        "created_at",
        "Ngày tạo DB",
        lambda item: safe_getattr(
            item,
            "created_at",
        ),
    ),
    "updated_at": ExportColumn(
        "updated_at",
        "Ngày cập nhật DB",
        lambda item: safe_getattr(
            item,
            "updated_at",
        ),
    ),
}


def normalize_cell(value: Any) -> Any:
    """
    Chuyển dữ liệu database thành kiểu có thể ghi được
    vào Excel, PDF và Google Sheets.
    """

    if value is None:
        return ""

    if isinstance(value, UUID):
        return str(value)

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, (dict, list)):
        return json.dumps(
            value,
            ensure_ascii=False,
        )

    return value


def resolve_columns(
    fields: list[str],
) -> list[ExportColumn]:
    unknown_fields = [field for field in fields if field not in EXPORT_COLUMNS]

    if unknown_fields:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": "Có trường export không hợp lệ.",
                "invalid_fields": unknown_fields,
            },
        )

    # Loại bỏ field bị gửi trùng nhưng vẫn giữ thứ tự.
    unique_fields = list(dict.fromkeys(fields))

    return [EXPORT_COLUMNS[field] for field in unique_fields]


def build_export_rows(
    items: list[Any],
    columns: list[ExportColumn],
) -> list[list[Any]]:
    return [
        [
            normalize_cell(
                column.getter(item),
            )
            for column in columns
        ]
        for item in items
    ]


def build_xlsx(
    headers: list[str],
    rows: list[list[Any]],
) -> bytes:
    """
    Dùng write_only để giảm RAM khi export nhiều dòng.
    """

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Listings")

    header_cells = []

    for header in headers:
        cell = WriteOnlyCell(
            worksheet,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="000000",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="9fc5e8",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        header_cells.append(cell)

    worksheet.freeze_panes = "A2"
    worksheet.append(header_cells)

    for row in rows:
        worksheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def register_pdf_fonts() -> None:
    """
    DejaVu Sans hỗ trợ đầy đủ tiếng Việt.
    Font được cài từ package fonts-dejavu-core.
    """

    registered_fonts = set(
        pdfmetrics.getRegisteredFontNames(),
    )

    if "DejaVuSans" not in registered_fonts:
        pdfmetrics.registerFont(
            TTFont(
                "DejaVuSans",
                ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            )
        )

    if "DejaVuSans-Bold" not in registered_fonts:
        pdfmetrics.registerFont(
            TTFont(
                "DejaVuSans-Bold",
                ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            )
        )


def build_pdf(
    title: str,
    headers: list[str],
    rows: list[list[Any]],
) -> bytes:
    register_pdf_fonts()

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.fontName = "DejaVuSans-Bold"
    title_style.fontSize = 14

    # Không dùng chung styles["BodyText"] cho header và body vì
    # ReportLab trả về cùng một object; sửa header sẽ làm body thành
    # chữ trắng và khiến dữ liệu trông như bị mất.
    cell_style = ParagraphStyle(
        "ExportCell",
        parent=styles["BodyText"],
        fontName="DejaVuSans",
        fontSize=6,
        leading=8,
        textColor=colors.black,
    )

    header_style = ParagraphStyle(
        "ExportHeader",
        parent=styles["BodyText"],
        fontName="DejaVuSans-Bold",
        fontSize=6,
        leading=8,
        textColor=colors.white,
    )

    pdf_rows = [
        [
            Paragraph(
                escape(str(header)),
                header_style,
            )
            for header in headers
        ]
    ]

    for row in rows:
        pdf_rows.append(
            [
                Paragraph(
                    escape(str(value)),
                    cell_style,
                )
                for value in row
            ]
        )

    available_width = landscape(A4)[0] - 20 * mm
    column_width = available_width / len(headers)

    table = Table(
        pdf_rows,
        repeatRows=1,
        colWidths=[column_width for _ in headers],
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1D4ED8"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.25,
                    colors.HexColor("#CBD5E1"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F8FAFC"),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    3,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    3,
                ),
            ]
        )
    )

    document.build(
        [
            Paragraph(title, title_style),
            table,
        ]
    )

    return buffer.getvalue()


def create_google_sheet(
    *,
    oauth_token_file: str,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
) -> tuple[str, str]:
    """
    Tạo một Google Spreadsheet mới và ghi dữ liệu
    thành từng batch để tránh request quá lớn.
    """

    credentials = get_google_credentials(
        oauth_token_file,
    )

    service = build(
        "sheets",
        "v4",
        credentials=credentials,
        cache_discovery=False,
    )

    spreadsheet = (
        service.spreadsheets()
        .create(
            body={
                "properties": {
                    "title": title,
                },
                "sheets": [
                    {
                        "properties": {
                            "title": "Listings",
                        }
                    }
                ],
            },
            fields=("spreadsheetId,spreadsheetUrl"),
        )
        .execute()
    )

    spreadsheet_id = spreadsheet["spreadsheetId"]

    # Header và dữ liệu được chia theo batch.
    all_rows = [headers, *rows]
    batch_size = 1_000

    for start_index in range(
        0,
        len(all_rows),
        batch_size,
    ):
        batch = all_rows[start_index : start_index + batch_size]

        # Google Sheets bắt đầu từ dòng 1.
        start_row = start_index + 1

        (
            service.spreadsheets()
            .values()
            .update(
                spreadsheetId=spreadsheet_id,
                range=f"Listings!A{start_row}",
                valueInputOption="RAW",
                body={
                    "values": batch,
                },
            )
            .execute()
        )

    spreadsheet_url = spreadsheet.get("spreadsheetUrl") or (
        f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit"
    )

    return spreadsheet_id, spreadsheet_url


async def export_marketplace_listings(
    *,
    session: AsyncSession,
    marketplace: str,
    payload: ListingExportRequest,
    settings: Settings,
) -> FileExportResult | GoogleSheetExportResult:
    columns = resolve_columns(payload.fields)

    filter_params = payload.filters.model_dump(
        exclude_none=True,
    )

    try:
        items, total = await list_listings(
            session,
            marketplace,
            page=1,
            page_size=EXPORT_MAX_ROWS,
            **filter_params,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=(status.HTTP_422_UNPROCESSABLE_ENTITY),
            detail=str(exc),
        ) from exc

    if total > EXPORT_MAX_ROWS:
        raise HTTPException(
            status_code=(status.HTTP_422_UNPROCESSABLE_ENTITY),
            detail=(
                f"Kết quả có {total} dòng. "
                f"Giới hạn export là "
                f"{EXPORT_MAX_ROWS} dòng. "
                "Hãy lọc bớt dữ liệu."
            ),
        )

    if payload.format == "pdf":
        if total > PDF_MAX_ROWS:
            raise HTTPException(
                status_code=(status.HTTP_422_UNPROCESSABLE_ENTITY),
                detail=(f"PDF chỉ hỗ trợ tối đa {PDF_MAX_ROWS} dòng."),
            )

        if len(columns) > PDF_MAX_COLUMNS:
            raise HTTPException(
                status_code=(status.HTTP_422_UNPROCESSABLE_ENTITY),
                detail=(f"PDF chỉ hỗ trợ tối đa {PDF_MAX_COLUMNS} cột."),
            )

    headers = [column.key for column in columns]

    rows = build_export_rows(
        items,
        columns,
    )

    now_text = datetime.now().strftime("%Y%m%d_%H%M%S")

    export_title = f"HQA {marketplace.upper()} Listings {now_text}"

    # ==================================================
    # 1. EXPORT EXCEL
    # ==================================================

    if payload.format == "xlsx":
        content = await to_thread.run_sync(
            build_xlsx,
            # Hai đối số này là positional arguments,
            # nên run_sync có thể chuyển trực tiếp.
            headers,
            rows,
        )

        return FileExportResult(
            kind="file",
            content=content,
            filename=(f"hqa_{marketplace}_{now_text}.xlsx"),
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            exported_rows=total,
        )

    # ==================================================
    # 2. EXPORT PDF
    # ==================================================

    if payload.format == "pdf":
        content = await to_thread.run_sync(
            build_pdf,
            # Các đối số positional có thể truyền
            # trực tiếp qua run_sync.
            export_title,
            headers,
            rows,
        )

        return FileExportResult(
            kind="file",
            content=content,
            filename=(f"hqa_{marketplace}_{now_text}.pdf"),
            media_type="application/pdf",
            exported_rows=total,
        )

    # ==================================================
    # 3. EXPORT GOOGLE SHEETS
    # ==================================================

    # Hàm create_google_sheet() dùng keyword-only
    # arguments nên không truyền trực tiếp các keyword
    # này vào to_thread.run_sync().
    create_sheet_job = partial(
        create_google_sheet,
        oauth_token_file=(settings.google_oauth_token_file),
        title=export_title,
        headers=headers,
        rows=rows,
    )

    # Worker thread sẽ gọi create_sheet_job().
    # create_sheet_job() sau đó gọi create_google_sheet()
    # với toàn bộ tham số đã được partial gắn sẵn.
    try:
        spreadsheet_id, spreadsheet_url = await to_thread.run_sync(
            create_sheet_job,
        )
    except FileNotFoundError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không tìm thấy Google OAuth token trên server.",
        ) from exc
    except RuntimeError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=str(exc),
        ) from exc
    except HttpError as exc:
        google_status = getattr(exc.resp, "status", None)
        if google_status in {401, 403}:
            detail = (
                "Google OAuth token không có quyền ghi Google Sheets. "
                "Hãy tạo lại token với scope spreadsheets (không phải readonly)."
            )
        else:
            detail = "Google Sheets API không thể tạo spreadsheet."
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=detail,
        ) from exc

    return GoogleSheetExportResult(
        kind="google_sheets",
        spreadsheet_id=spreadsheet_id,
        spreadsheet_url=spreadsheet_url,
        exported_rows=total,
    )
